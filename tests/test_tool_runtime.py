import asyncio
from pathlib import Path
import subprocess
import tempfile
import unittest
import zipfile

import openpyxl

from MAS_DAG.mas_runtime import GenerationResult, run_candidate_graph, run_candidate_graph_async
from MAS_DAG.tool_runtime import EpisodeWorkspace, execute_tool


class ScriptedBackend:
    def __init__(self, responses):
        self.responses = list(responses)

    def count_tokens(self, text):
        return len(text.split())

    def generate(self, messages, *, max_new_tokens=None):
        text = self.responses.pop(0)
        return GenerationResult(
            text=text,
            input_tokens=sum(self.count_tokens(x["content"]) for x in messages),
            output_tokens=self.count_tokens(text),
            latency_seconds=0.01,
        )


class AsyncScriptedBackend(ScriptedBackend):
    async def generate(self, messages, *, max_new_tokens=None):
        await asyncio.sleep(0)
        return super().generate(messages, max_new_tokens=max_new_tokens)


class ToolRuntimeTests(unittest.TestCase):
    def test_episode_workspace_and_spreadsheet_tool(self):
        with tempfile.TemporaryDirectory() as source_dir:
            source = Path(source_dir) / "table.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(["name", "year"])
            sheet.append(["oldest", 2009])
            workbook.save(source)
            with EpisodeWorkspace({"attachment_source": str(source)}) as workspace:
                listing = execute_tool("workspace_list_files", {}, workspace)
                self.assertEqual(listing["files"], ["attachments/table.xlsx"])
                result = execute_tool(
                    "workspace_read_spreadsheet", {"path": "table.xlsx"}, workspace
                )
                self.assertEqual(result["rows"][1], ["oldest", 2009])
                with self.assertRaisesRegex(ValueError, "escapes"):
                    workspace.resolve("../../etc/passwd")

    def test_document_archive_and_image_tools(self):
        with tempfile.TemporaryDirectory() as source_dir:
            root = Path(source_dir)
            docx = root / "sample.docx"
            with zipfile.ZipFile(docx, "w") as archive:
                archive.writestr(
                    "word/document.xml",
                    '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Contract evidence</w:t></w:r></w:p></w:body></w:document>',
                )
            with EpisodeWorkspace({"attachment_source": str(docx)}) as workspace:
                result = execute_tool(
                    "workspace_read_document", {"path": "sample.docx"}, workspace
                )
                self.assertIn("Contract evidence", result["text"])

            archive_path = root / "bundle.zip"
            with zipfile.ZipFile(archive_path, "w") as archive:
                archive.writestr("notes/evidence.txt", "verified value: 42")
            with EpisodeWorkspace({"attachment_source": str(archive_path)}) as workspace:
                listing = execute_tool(
                    "workspace_list_archive", {"path": "bundle.zip"}, workspace
                )
                self.assertEqual(listing["members"][0]["name"], "notes/evidence.txt")
                result = execute_tool(
                    "workspace_read_archive_member",
                    {"path": "bundle.zip", "member": "notes/evidence.txt"},
                    workspace,
                )
                self.assertIn("42", result["text"])
                extracted = execute_tool(
                    "workspace_extract_archive_member",
                    {"path": "bundle.zip", "member": "notes/evidence.txt"},
                    workspace,
                )
                self.assertEqual(
                    (workspace.root / extracted["path"]).read_text(), "verified value: 42"
                )
                with self.assertRaisesRegex(ValueError, "unsafe"):
                    execute_tool(
                        "workspace_read_archive_member",
                        {"path": "bundle.zip", "member": "../escape.txt"},
                        workspace,
                    )

            from PIL import Image

            image_path = root / "sample.png"
            Image.new("RGB", (17, 23), "white").save(image_path)
            with EpisodeWorkspace({"attachment_source": str(image_path)}) as workspace:
                result = execute_tool(
                    "workspace_inspect_file", {"path": "sample.png"}, workspace
                )
                self.assertEqual(result["image"]["width"], 17)
                self.assertEqual(result["image"]["height"], 23)

    def test_structured_json_pdb_and_spreadsheet_style_tools(self):
        with tempfile.TemporaryDirectory() as source_dir:
            root = Path(source_dir)
            jsonld = root / "record.jsonld"
            jsonld.write_text('{"@type":"Person","name":"Ada"}', encoding="utf-8")
            with EpisodeWorkspace({"attachment_source": str(jsonld)}) as workspace:
                result = execute_tool("workspace_read_json", {"path": jsonld.name}, workspace)
                self.assertIn('"Ada"', result["content"])

            pdb = root / "sample.pdb"
            pdb.write_text(
                "ATOM      1  N   ALA A   1      11.104  13.207   9.900  1.00 20.00           N\n"
                "ATOM      2  N   GLY A   2      12.104  13.207   9.900  1.00 20.00           N\n",
                encoding="utf-8",
            )
            with EpisodeWorkspace({"attachment_source": str(pdb)}) as workspace:
                result = execute_tool("workspace_analyze_pdb", {"path": pdb.name}, workspace)
                self.assertEqual(result["chains"]["A"]["sequence"], "AG")

            xlsx = root / "styled.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet["A1"] = "START"
            sheet["A1"].fill = openpyxl.styles.PatternFill("solid", fgColor="00FF00")
            workbook.save(xlsx)
            with EpisodeWorkspace({"attachment_source": str(xlsx)}) as workspace:
                result = execute_tool(
                    "workspace_inspect_spreadsheet", {"path": xlsx.name}, workspace
                )
                self.assertEqual(result["cells"][0]["coordinate"], "A1")
                self.assertEqual(result["cells"][0]["fill_fg"], "0000FF00")
                grid = execute_tool(
                    "workspace_analyze_spreadsheet_grid",
                    {"path": xlsx.name, "fill_rgb": "00FF00"},
                    workspace,
                )
                self.assertEqual(grid["cell_count"], 1)
                self.assertTrue(grid["hamiltonian_cycle_ruled_out"])

    def test_sync_graph_runs_tool_loop_and_records_trace(self):
        nodes = [
            {
                "id": "s",
                "role": "calculator",
                "role_brief": "Calculate with tools.",
                "tools": ["python_calculate"],
            },
            {"id": "f", "role": "finalizer", "role_brief": "Return the answer."},
        ]
        graph = {
            "mask": [0, 0],
            "edge_weight": [[0, 1], [0, 0]],
        }
        backend = ScriptedBackend([
            'TOOL_CALL: {"name":"python_calculate","arguments":{"code":"print(6 * 7)"}}',
            "NODE_RESULT: Python returned 42.",
            "FINAL_ANSWER: 42",
        ])
        result = run_candidate_graph(
            task="Compute six times seven.",
            reference_answer="42",
            nodes=nodes,
            graph=graph,
            finalizer_id="f",
            backend=backend,
            evaluator="gaia",
            enable_tools=True,
        )
        self.assertEqual(result["accuracy"], 1.0)
        self.assertEqual(result["total_tool_calls"], 1)
        self.assertEqual(result["node_tool_traces"][0][0]["status"], "ok")
        self.assertIn("42", result["node_tool_traces"][0][0]["result"]["stdout"])

    def test_media_probe_and_video_frame_extraction(self):
        try:
            import imageio_ffmpeg
        except ImportError:
            self.skipTest("GAIA media extras are not installed")
        with tempfile.TemporaryDirectory() as source_dir:
            video = Path(source_dir) / "sample.mp4"
            subprocess.run(
                [
                    imageio_ffmpeg.get_ffmpeg_exe(), "-hide_banner", "-loglevel", "error",
                    "-f", "lavfi", "-i", "color=c=blue:s=160x120:d=2:r=2",
                    "-pix_fmt", "yuv420p", str(video),
                ],
                check=True,
            )
            with EpisodeWorkspace({"attachment_source": str(video)}) as workspace:
                probe = execute_tool("workspace_probe_media", {"path": video.name}, workspace)
                self.assertEqual(probe["streams"][0]["width"], 160)
                frames = execute_tool(
                    "workspace_extract_video_frames",
                    {"path": video.name, "interval_seconds": 1, "max_frames": 2},
                    workspace,
                )
                self.assertEqual(len(frames["frames"]), 2)
                self.assertTrue(all((workspace.root / item).is_file() for item in frames["frames"]))

    def test_async_graph_runs_tool_loop(self):
        nodes = [
            {
                "id": "s",
                "role": "calculator",
                "role_brief": "Calculate with tools.",
                "tools": ["python_calculate"],
            },
            {"id": "f", "role": "finalizer", "role_brief": "Return the answer."},
        ]
        graph = {"mask": [0, 0], "edge_weight": [[0, 1], [0, 0]]}
        backend = AsyncScriptedBackend([
            'TOOL_CALL: {"name":"python_calculate","arguments":{"code":"print(21 * 2)"}}',
            "NODE_RESULT: The calculation is 42.",
            "FINAL_ANSWER: 42",
        ])
        result = asyncio.run(run_candidate_graph_async(
            task="Compute.",
            reference_answer="42",
            nodes=nodes,
            graph=graph,
            finalizer_id="f",
            backend=backend,
            evaluator="gaia",
            enable_tools=True,
        ))
        self.assertEqual(result["accuracy"], 1.0)
        self.assertEqual(result["total_tool_calls"], 1)


if __name__ == "__main__":
    unittest.main()
